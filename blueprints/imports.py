"""Blueprint imports : import TSV/CSV/vCard et export TSV/vCard des contacts.

Endpoints : imports.index (page + traitement import), imports.export_vcard,
imports.export_contacts.
"""
import csv
import io
import os
import time
import uuid
import unicodedata
from datetime import datetime

from flask import (Blueprint, render_template, request, redirect, url_for,
                   flash, Response)
from flask_login import login_required, current_user

from models import db, Contact, Liste, CustomFieldDefinition, ImportMapping
from vcard_converter import extract_vcard_data, get_vcards, MULTI_VALUE_SEP
from helpers import admin_required, slugify_key, listes_sorted
import fields as fields_registry

bp = Blueprint('imports', __name__)

# Dossier temporaire des fichiers en cours d'import (mappage en 2 temps)
_IMPORT_DIR = 'data/imports'

# Alias de noms de colonnes → clé de champ (auto-suggestion du mapping)
_MAPPING_ALIASES = {
    'courriel': 'email', 'mail': 'email', 'e mail': 'email', 'adresse mail': 'email', 'adresse email': 'email',
    'first name': 'prenom', 'firstname': 'prenom', 'given name': 'prenom',
    'nom de famille': 'nom', 'last name': 'nom', 'lastname': 'nom', 'surname': 'nom',
    'tel': 'telephone', 'phone': 'telephone', 'portable': 'telephone', 'mobile': 'telephone',
    'gsm': 'telephone', 'tel portable': 'telephone', 'numero': 'telephone',
    'commune': 'adresse_ville', 'city': 'adresse_ville',
    'code postal': 'adresse_cp', 'cp': 'adresse_cp', 'zip': 'adresse_cp', 'zip code': 'adresse_cp',
    'adresse': 'adresse_rue', 'rue': 'adresse_rue', 'street': 'adresse_rue', 'voie': 'adresse_rue',
    'complement': 'adresse_complement', 'complement adresse': 'adresse_complement',
    'departement': 'adresse_region', 'region': 'adresse_region',
    'pays': 'adresse_pays', 'country': 'adresse_pays',
    'organization': 'organisation', 'societe': 'organisation', 'entreprise': 'organisation', 'company': 'organisation',
    'note': 'notes', 'remarque': 'notes', 'remarques': 'notes', 'commentaire': 'notes',
    'liste': 'listes', 'categorie': 'listes', 'categories': 'listes', 'groupe': 'listes', 'groupes': 'listes',
}


# === Helpers d'import ===

def _parse_liste_names(raw):
    """Parse les noms de listes depuis une chaîne (vCard: 'Catégories', TSV: 'Listes')."""
    if not raw:
        return []
    raw = raw.strip().replace('[', '').replace(']', '').replace("'", '')
    for sep in [MULTI_VALUE_SEP.strip(), ',']:
        if sep in raw:
            return [c.strip() for c in raw.split(sep) if c.strip()]
    return [raw.strip()] if raw.strip() else []


def _extract_fields_from_row(row):
    """Extrait les champs normalisés depuis un dict (TSV ou vCard).

    Accepte les colonnes 'Listes', 'Catégories' ou 'Categories' pour les listes
    (rétrocompatibilité vCard et TSV).
    """
    # Email
    email_val = (
        row.get('Email', '') or
        row.get('email', '') or
        row.get('Email_Home', '') or
        row.get('Email_Work', '') or
        row.get('Email_Autre', '')
    ).strip()
    if MULTI_VALUE_SEP.strip() in email_val:
        email_val = email_val.split(MULTI_VALUE_SEP.strip())[0].strip()

    # Nom / Prénom
    nom = ''
    prenom = ''
    nom_prenom = row.get('Nom, Prénom', '')
    if nom_prenom:
        parts = nom_prenom.split(',', 1)
        nom = parts[0].strip()
        prenom = parts[1].strip() if len(parts) > 1 else ''

    if not nom:
        nom = row.get('Nom', row.get('nom', '')).strip()
    if not prenom:
        prenom = row.get('Prenom', row.get('prenom', row.get('Prénom', ''))).strip()

    if not nom and not prenom:
        fn = row.get('Nom Complet', '').strip()
        if fn:
            parts = fn.rsplit(' ', 1)
            if len(parts) == 2:
                prenom, nom = parts
            else:
                nom = fn

    # Téléphone
    telephone = (
        row.get('Tel_Cell', '') or
        row.get('Tel_Home', '') or
        row.get('Tel_Work', '') or
        row.get('telephone', '') or
        row.get('Tel', '')
    ).strip()
    if MULTI_VALUE_SEP.strip() in telephone:
        telephone = telephone.split(MULTI_VALUE_SEP.strip())[0].strip()

    # Listes (accepte 'Listes', 'Catégories', 'Categories' pour rétrocompatibilité)
    listes_raw = (
        row.get('Listes', '') or
        row.get('Catégories', '') or
        row.get('Categories', '')
    )
    listes = _parse_liste_names(listes_raw)

    return {
        'uid': row.get('UID', row.get('uid', '')).strip(),
        'email': email_val,
        'nom': nom,
        'prenom': prenom,
        'genre': row.get('Genre', row.get('genre', '')).strip(),
        'titre': row.get('Titre', row.get('titre', '')).strip(),
        'telephone': telephone,
        'organisation': row.get('Organisation', row.get('organisation', '')).strip(),
        'adresse_rue': row.get('Rue', row.get('adresse_rue', '')).strip(),
        'adresse_complement': row.get('Complement', row.get('adresse_complement', '')).strip(),
        'adresse_ville': row.get('Ville', row.get('adresse_ville', '')).strip(),
        'adresse_cp': row.get('CP', row.get('adresse_cp', '')).strip(),
        'adresse_region': row.get('Region', row.get('adresse_region', '')).strip(),
        'adresse_pays': row.get('Pays', row.get('adresse_pays', '')).strip(),
        'source': row.get('Source', row.get('source', '')).strip(),
        'notes': row.get('Note', row.get('Notes', row.get('notes', ''))).strip(),
        'listes': listes,
    }


def _get_or_create_listes(noms):
    """Retourne les objets Liste pour une liste de noms, en créant ceux qui n'existent pas."""
    listes = []
    for nom in noms:
        liste = Liste.query.filter_by(nom=nom).first()
        if not liste:
            liste = Liste(nom=nom)
            db.session.add(liste)
        listes.append(liste)
    return listes


def _detect_vcard_source(content):
    """Détecte la source d'un fichier vCard depuis son contenu (PRODID, format UID)."""
    content_lower = content.lower()
    if 'prodid' in content_lower:
        if 'roundcube' in content_lower:
            return 'Roundcube'
        if 'infomaniak' in content_lower:
            return 'Infomaniak'
        if 'proton' in content_lower:
            return 'Proton'
        if 'thunderbird' in content_lower or 'cardbook' in content_lower:
            return 'Thunderbird'
        if 'apple' in content_lower or 'addressbook' in content_lower:
            return 'Apple'
        if 'google' in content_lower:
            return 'Google'
    # Heuristiques sur le format UID
    if 'uid:proton-' in content_lower:
        return 'Proton'
    # UID Roundcube/SOGo : 32hex-16hex (pas de PRODID)
    import re
    if re.search(r'UID:[0-9A-F]{32}-[0-9A-F]{16}', content):
        return 'Roundcube'
    return 'vCard'


def _import_contact_from_row(row, update_existing=False, source='Import'):
    """
    Importe un contact depuis un dict (TSV ou vCard).

    Détection des doublons :
      1. Par UID (identité exacte, si présent dans le fichier importé)
      2. Par composite email + nom + prénom (même personne probable)
      3. Sinon → nouveau contact créé (même si l'email existe déjà)

    Retourne (contact, action) où action = 'created', 'updated', 'no_email' ou 'skipped'.
    """
    fields = _extract_fields_from_row(row)

    if not fields['email']:
        return None, 'no_email'

    existing = None

    # Priorité 1 : correspondance par UID
    if fields['uid']:
        existing = Contact.query.filter_by(uid=fields['uid'], is_deleted=False).first()

    # Priorité 2 : correspondance composite email + nom + prénom
    if not existing and fields['nom'] and fields['prenom']:
        existing = Contact.query.filter_by(
            email=fields['email'],
            nom=fields['nom'],
            prenom=fields['prenom'],
            is_deleted=False
        ).first()

    if existing and not update_existing:
        return None, 'skipped'

    if existing and update_existing:
        # Mettre à jour les champs non vides
        if fields['nom']:
            existing.nom = fields['nom']
        if fields['prenom']:
            existing.prenom = fields['prenom']
        if fields['genre']:
            existing.genre = fields['genre']
        if fields['titre']:
            existing.titre = fields['titre']
        if fields['telephone']:
            existing.telephone = fields['telephone']
        if fields['organisation']:
            existing.organisation = fields['organisation']
        if fields['adresse_rue']:
            existing.adresse_rue = fields['adresse_rue']
        if fields['adresse_complement']:
            existing.adresse_complement = fields['adresse_complement']
        if fields['adresse_ville']:
            existing.adresse_ville = fields['adresse_ville']
        if fields['adresse_cp']:
            existing.adresse_cp = fields['adresse_cp']
        if fields['adresse_region']:
            existing.adresse_region = fields['adresse_region']
        if fields['adresse_pays']:
            existing.adresse_pays = fields['adresse_pays']
        if fields['notes']:
            existing.notes = fields['notes']

        # Remplacement des listes par celles de l'import
        if fields['listes']:
            existing.listes = _get_or_create_listes(fields['listes'])

        return existing, 'updated'

    # Nouveau contact (même si l'email existe déjà chez un autre contact)
    kwargs = dict(
        nom=fields['nom'],
        prenom=fields['prenom'],
        genre=fields['genre'],
        titre=fields['titre'],
        email=fields['email'],
        telephone=fields['telephone'],
        organisation=fields['organisation'],
        adresse_rue=fields['adresse_rue'],
        adresse_complement=fields['adresse_complement'],
        adresse_ville=fields['adresse_ville'],
        adresse_cp=fields['adresse_cp'],
        adresse_region=fields['adresse_region'],
        adresse_pays=fields['adresse_pays'],
        notes=fields['notes'],
        source=fields.get('source') or source
    )
    # Préserver le UID d'origine (Roundcube, Proton, etc.) s'il est fourni
    if fields['uid']:
        kwargs['uid'] = fields['uid']
    contact = Contact(**kwargs)
    contact.listes = _get_or_create_listes(fields['listes'])

    return contact, 'created'


# === Import v2 : mapping registre-driven (Excel/CSV) ===

def _norm(s):
    """Normalise pour matcher : minuscule, sans accents, alphanum, espaces compactés."""
    s = unicodedata.normalize('NFKD', str(s or '')).encode('ascii', 'ignore').decode()
    s = ''.join(c if c.isalnum() else ' ' for c in s.lower())
    return ' '.join(s.split())


def _import_targets():
    """Champs mappables : registre (cœur + perso, éditables) + synthétiques listes/uid."""
    targets = []
    for f in fields_registry.contact_fields(include_custom=True):
        if f.editable:   # 'source' est editable=False → jamais proposé à l'import
            targets.append({'key': f.key, 'label': f.label, 'group': f.group,
                            'custom': (f.source == 'custom')})
    targets.append({'key': 'listes', 'label': 'Listes (séparées par ;)', 'group': 'listes', 'custom': False})
    targets.append({'key': 'uid', 'label': 'UID (dédoublonnage)', 'group': 'systeme', 'custom': False})
    return targets


def _suggest_mapping(headers, targets):
    """Devine {header: clé} par nom (clé/label du registre, puis alias)."""
    by_norm = {}
    for t in targets:
        by_norm.setdefault(_norm(t['label']), t['key'])
        by_norm.setdefault(_norm(t['key']), t['key'])
    out = {}
    for h in headers:
        n = _norm(h)
        out[h] = by_norm.get(n) or _MAPPING_ALIASES.get(n) or ''
    return out


def _sweep_stale_imports(max_age_h=6):
    """Best-effort : purge les fichiers temporaires d'import abandonnés (aperçu
    sans validation, session fermée…). Silencieux."""
    try:
        cutoff = time.time() - max_age_h * 3600
        for name in os.listdir(_IMPORT_DIR):
            p = os.path.join(_IMPORT_DIR, name)
            try:
                if os.path.isfile(p) and os.path.getmtime(p) < cutoff:
                    os.remove(p)
            except OSError:
                pass
    except OSError:
        pass


def _cell_str(v):
    """Cellule openpyxl → chaîne « propre ». Excel stocke tout nombre en flottant :
    un entier ressort en float (34700.0) → on retire le « .0 » parasite (codes postaux,
    identifiants…) tout en préservant les vraies décimales (12.5)."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_tabular(path, filename):
    """(headers, rows) depuis .xlsx / .csv / .tsv. rows = list de {header: str}."""
    fn = (filename or '').lower()
    if fn.endswith('.xlsx') or fn.endswith('.xlsm'):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        try:
            head = next(it)
        except StopIteration:
            wb.close(); return [], []
        # Ne garder que les colonnes à EN-TÊTE non vide : openpyxl gonfle souvent la
        # dimension d'une vraie feuille avec des colonnes vides en fin → sinon on
        # affiche des dizaines/centaines de « colonne N » fantômes. On mémorise l'index
        # d'origine de chaque colonne réelle pour lire la bonne cellule dans les lignes.
        cols = [(i, _cell_str(h)) for i, h in enumerate(head)
                if h is not None and _cell_str(h).strip()]
        headers = [name for _, name in cols]
        rows = []
        for r in it:
            d = {name: ('' if (idx >= len(r) or r[idx] is None) else _cell_str(r[idx]))
                 for idx, name in cols}
            if any(d.values()):
                rows.append(d)
        wb.close()
        return headers, rows
    # CSV / TSV : utf-8 (BOM géré) sinon cp1252 (encodage réel des exports Excel FR).
    # cp1252 mappe les 256 octets → jamais d'échec, et rend correctement é/è/à/ç…
    # (on évite la détection statistique, peu fiable sur de petits fichiers).
    with open(path, 'rb') as fp:
        raw = fp.read()
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = raw.decode('cp1252', errors='replace')
    first = text.split('\n', 1)[0]
    delim = '\t' if '\t' in first else (';' if first.count(';') > first.count(',') else ',')
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    headers = [(h or '').strip() for h in (reader.fieldnames or [])]
    rows = []
    for r in reader:
        d = {(k or '').strip(): ('' if v is None else str(v).strip()) for k, v in r.items()}
        if any(d.values()):
            rows.append(d)
    return headers, rows


def _apply_mapping(row, mapping):
    """{clé_champ: valeur} depuis {header: valeur} + {header: clé}. 1re valeur non vide gagne."""
    mapped = {}
    for header, key in mapping.items():
        if not key:
            continue
        val = (row.get(header) or '').strip()
        if val and not mapped.get(key):
            mapped[key] = val
    return mapped


def _key_sets():
    """(clés colonnes, clés perso) éditables du registre."""
    col_keys, custom_keys = set(), set()
    for f in fields_registry.contact_fields(include_custom=True):
        if not f.editable:
            continue
        (custom_keys if f.source == 'custom' else col_keys).add(f.key)
    return col_keys, custom_keys


def _custom_types():
    """{clé_perso: type} — pour coercition des valeurs importées selon le type du champ."""
    return {f.key: f.type for f in fields_registry.contact_fields(include_custom=True)
            if f.source == 'custom'}


# Jetons interprétés comme « vrai » pour un champ case à cocher (le reste = faux).
_TRUTHY = {'1', 'true', 'vrai', 'oui', 'yes', 'y', 'x', 'o', 'v'}


def _norm_date(v):
    """Normalise une date en ISO YYYY-MM-DD si reconnue, sinon renvoie la valeur brute.
    Gère notamment le datetime openpyxl (« 2026-08-06 00:00:00 ») et le format FR."""
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(v, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return v


def _coerce_custom(ftype, val):
    """Coerce une valeur importée selon le type du champ perso.
    Renvoie None quand le champ doit rester VIDE (absent de custom_fields)."""
    v = (val or '').strip()
    if not v:
        return None
    if ftype == 'checkbox':
        # « FALSE / non / 0 » → non coché (None) ; « TRUE / oui / 1 / x » → '1'
        return '1' if v.lower() in _TRUTHY else None
    if ftype == 'date':
        return _norm_date(v)
    return v


def _coerce_column(key, val):
    """Normalisation légère des champs COLONNE à l'import.
    Code postal FR : un CP à 1-4 chiffres a perdu son/ses zéro(s) de tête au passage
    par un nombre Excel (06123 → 6123) → on recomplète à 5 chiffres."""
    if key == 'adresse_cp' and val.isdigit() and 0 < len(val) < 5:
        return val.zfill(5)
    return val


def _genre_from_civilite(civ):
    """Genre grammatical déduit de la civilité (pour l'accord des mailings).
    Renvoie None si ambigu (Mx/Autre…) → on garde le défaut (Inclusif)."""
    c = (civ or '').strip().lower().rstrip('.')
    if c in ('m', 'mr', 'monsieur'):
        return 'Masculin'
    if c in ('mme', 'mlle', 'madame', 'mademoiselle'):
        return 'Féminin'
    return None


def _import_mapped(mapped, col_keys, custom_keys, custom_types, update_existing, source, extra_listes):
    """Importe une row MAPPÉE (keyée par clé de champ). Sans email = ACCEPTÉ.
    Dédup : UID puis composite email+nom+prénom (si email). Retourne (contact, action)."""
    email = (mapped.get('email') or '').strip()
    nom = (mapped.get('nom') or '').strip()
    prenom = (mapped.get('prenom') or '').strip()
    uid = (mapped.get('uid') or '').strip()

    existing = None
    if uid:
        existing = Contact.query.filter_by(uid=uid, is_deleted=False).first()
    if not existing and nom and prenom:
        # Dédoublonnage par nom+prénom ; l'email affine quand il est présent
        # (familles au même email distinguées par le prénom). Sans email — cas
        # fréquent (fichier « maires » = 0 email) — nom+prénom suffit à retrouver.
        q = Contact.query.filter_by(nom=nom, prenom=prenom, is_deleted=False)
        if email:
            q = q.filter_by(email=email)
        existing = q.first()

    if existing and not update_existing:
        return existing, 'skipped'

    listes_names = sorted(set(_parse_liste_names(mapped.get('listes', '')) + list(extra_listes)))

    def _write(contact):
        for key, val in mapped.items():
            if not val or key in ('listes', 'uid'):
                continue
            if key in col_keys:
                setattr(contact, key, _coerce_column(key, val))
            elif key in custom_keys:
                coerced = _coerce_custom(custom_types.get(key, 'text'), val)
                cf = dict(contact.custom_fields or {})
                if coerced is None:
                    cf.pop(key, None)   # valeur « fausse »/vide → champ laissé absent
                else:
                    cf[key] = coerced
                contact.custom_fields = cf
        # Accord de genre auto : si la civilité est fournie mais PAS de colonne Genre
        # explicite, on déduit le genre grammatical (M./Monsieur→Masculin, Mme→Féminin).
        # Une colonne Genre du fichier reste prioritaire ; ambiguïté → défaut (Inclusif).
        civ_in = (mapped.get('civilite') or '').strip()
        if civ_in and not (mapped.get('genre') or '').strip():
            derived = _genre_from_civilite(civ_in)
            if derived:
                contact.genre = derived
        if uid and not contact.uid:
            contact.uid = uid
        if listes_names:
            objs = _get_or_create_listes(listes_names)
            if update_existing and existing is contact:
                contact.listes = objs
            else:
                for l in objs:
                    if l not in contact.listes:
                        contact.listes.append(l)

    if existing and update_existing:
        _write(existing)
        return existing, 'updated'

    # Nouveau : colonnes NOT NULL initialisées à '' (contact « à compléter » autorisé)
    contact = Contact(nom='', prenom='', email='', source=(mapped.get('source') or source))
    db.session.add(contact)   # en session AVANT d'attacher les listes (sinon l'assoc n'est pas prise)
    _write(contact)
    return contact, 'created'


def _dry_run(rows, mapping, col_keys, custom_keys, custom_types, update_existing, extra_listes):
    """Compte created/updated/skipped SANS écrire (rollback à la fin)."""
    counts = {'created': 0, 'updated': 0, 'skipped': 0}
    sample = []
    for i, row in enumerate(rows):
        mapped = _apply_mapping(row, mapping)
        _c, action = _import_mapped(mapped, col_keys, custom_keys, custom_types, update_existing, 'preview', extra_listes)
        counts[action] = counts.get(action, 0) + 1
        no_email = not (mapped.get('email') or '').strip()
        if no_email and action != 'skipped':
            counts['no_email'] = counts.get('no_email', 0) + 1
        if i < 6:
            sample.append({'mapped': mapped, 'action': action, 'no_email': no_email})
    db.session.rollback()   # défausse toute mutation — rien n'est persisté
    return counts, sample


def _run_import(rows, mapping, col_keys, custom_keys, custom_types, update_existing, extra_listes, user_id):
    # Auto-backup AVANT toute écriture : un import (surtout création de champs perso /
    # mise à jour en masse) est difficilement réversible → snapshot cohérent pour
    # pouvoir revenir en arrière. Best-effort (ne bloque pas l'import) mais loggé.
    from datetime import datetime as _dt
    from helpers import backup_database
    counts = {'created': 0, 'updated': 0, 'skipped': 0, 'no_email': 0, 'backup': None}
    try:
        counts['backup'] = backup_database(f"data/backups/pre-import-{_dt.now().strftime('%Y%m%d-%H%M%S')}.db")
    except Exception:
        import logging
        logging.exception('auto-backup pré-import échoué (import poursuivi)')
    for row in rows:
        mapped = _apply_mapping(row, mapping)
        contact, action = _import_mapped(mapped, col_keys, custom_keys, custom_types, update_existing, 'Import', extra_listes)
        if action == 'created':
            contact.created_by_id = user_id
            db.session.add(contact)
        counts[action] += 1
        if not (mapped.get('email') or '').strip() and action != 'skipped':
            counts['no_email'] += 1
    db.session.commit()
    return counts


def _mapping_from_form(form):
    """Reconstruit {header: clé} depuis les paires cachées hdr/map (ordre préservé)."""
    return dict(zip(form.getlist('hdr'), form.getlist('map')))


# --- Mappings enregistrés (réutilisables) ---

def _saved_mappings():
    return ImportMapping.query.order_by(ImportMapping.name).all()


def _apply_saved(headers, saved):
    """{header: clé} pour les en-têtes du fichier couverts par un mapping enregistré,
    par correspondance de NOM normalisé (robuste à l'ordre / colonnes manquantes)."""
    lut = {_norm(h): k for h, k in (saved.mapping or {}).items() if k}
    out = {}
    for h in headers:
        k = lut.get(_norm(h))
        if k:
            out[h] = k
    return out


def _best_saved(headers, saved_list):
    """Mapping enregistré couvrant le plus d'en-têtes du fichier (par nom normalisé).
    Renvoie (ImportMapping|None, nb_correspondances)."""
    hset = {_norm(h) for h in headers}
    best, best_n = None, 0
    for sm in saved_list:
        keys = {_norm(h) for h, k in (sm.mapping or {}).items() if k}
        n = len(keys & hset)
        if n > best_n:
            best, best_n = sm, n
    return best, best_n


def _render_mapping(token, ext, filename, headers, rows, mapping, list_id,
                    new_list_name, update_existing, saved_mappings,
                    applied_mapping=None, previewed=False, counts=None):
    """Rendu unique de l'écran de mapping (upload, aperçu, save, apply)."""
    return render_template('import_mapping.html',
                           token=token, ext=ext, filename=filename,
                           headers=headers, sample_rows=rows[:5], nrows=len(rows),
                           targets=_import_targets(), mapping=mapping,
                           listes=_active_listes(), list_id=list_id,
                           new_list_name=new_list_name, update_existing=update_existing,
                           dedup_ok=_dedup_ok(mapping), new_field_types=_NEW_FIELD_TYPES,
                           saved_mappings=saved_mappings, applied_mapping=applied_mapping,
                           previewed=previewed, counts=counts,
                           field_map=fields_registry.field_map())


# Types autorisés à la création d'un champ perso depuis l'import (MVP-2).
_NEW_FIELD_TYPES = (
    ('text', 'Texte'), ('textarea', 'Texte long'), ('number', 'Nombre'),
    ('date', 'Date'), ('checkbox', 'Case à cocher'), ('select', 'Liste déroulante'),
)
_NEW_FIELD_TYPE_KEYS = {k for k, _ in _NEW_FIELD_TYPES}


def _create_import_fields(mapping, form):
    """Crée les champs perso demandés à l'import (colonnes mappées sur « __new__ »).

    Réécrit `mapping` en place (header → clé du nouveau champ) et renvoie la liste
    des champs créés [(libellé, clé)]. Idempotent : une clé déjà existante est
    réutilisée (pas de doublon si on prévisualise puis importe). Persiste aussitôt
    (définitions légères, gérables ensuite dans Paramètres › Champs personnalisés)."""
    newlabels = dict(zip(form.getlist('hdr'), form.getlist('newlabel')))
    newtypes = dict(zip(form.getlist('hdr'), form.getlist('newtype')))
    created = []
    for hdr, key in list(mapping.items()):
        if key != '__new__':
            continue
        label = (newlabels.get(hdr) or hdr or '').strip()
        ftype = newtypes.get(hdr) if newtypes.get(hdr) in _NEW_FIELD_TYPE_KEYS else 'text'
        slug = slugify_key(label)
        if not slug or slug in fields_registry.RESERVED_KEYS:
            mapping[hdr] = ''   # clé impossible/réservée → colonne non importée
            continue
        cf = CustomFieldDefinition.query.filter_by(key=slug).first()
        if not cf:
            max_ordre = db.session.query(db.func.max(CustomFieldDefinition.ordre)).scalar() or 0
            db.session.add(CustomFieldDefinition(
                key=slug, display_name=label, type=ftype, ordre=max_ordre + 1))
            created.append((label, slug))
        mapping[hdr] = slug
    if created:
        db.session.commit()
    return created


def _dedup_ok(mapping):
    """Une clé de dédoublonnage est-elle mappée ? (UID, ou Nom+Prénom).
    Sinon un ré-import ne peut PAS retrouver les contacts existants → il crée des doublons.
    L'email n'est pas requis (souvent absent) : il ne fait qu'affiner le nom+prénom."""
    keys = {v for v in mapping.values() if v}
    return ('uid' in keys) or ({'nom', 'prenom'} <= keys)


def _active_listes():
    return listes_sorted(is_archived=False)


def _resolve_target_list(list_id, new_list_name):
    """Résout la liste de destination. Renvoie (noms_a_ajouter, id_pour_redirection).

    Une nouvelle liste (nom saisi) est créée à la volée (flush → id dispo) ; en
    dry-run le rollback l'annule, en import réel le commit la persiste."""
    new_list_name = (new_list_name or '').strip()
    if new_list_name:
        l = Liste.query.filter_by(nom=new_list_name).first()
        if not l:
            l = Liste(nom=new_list_name)
            db.session.add(l)
            db.session.flush()
        return [l.nom], l.id
    if list_id:
        l = Liste.query.get(int(list_id))
        if l:
            return [l.nom], l.id
    return [], None


# === Routes ===

def _import_vcard_direct(file):
    """Import direct d'un fichier vCard (schéma standard → pas de mapping)."""
    import tempfile
    update_existing = request.form.get('update_existing') == 'on'
    content = file.read()
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.vcf', mode='wb')
    tmp.write(content)
    tmp.close()
    created = updated = skipped = no_email = 0
    try:
        source = _detect_vcard_source(content.decode('utf-8', errors='replace'))
        for vcard in get_vcards(tmp.name):
            row = extract_vcard_data(vcard, tmp.name)
            contact, action = _import_contact_from_row(row, update_existing=update_existing, source=source)
            if action == 'created':
                contact.created_by_id = current_user.id
                db.session.add(contact)
                created += 1
            elif action == 'updated':
                updated += 1
            elif action == 'skipped':
                skipped += 1
            else:
                no_email += 1
        db.session.commit()
        parts = []
        if created:
            parts.append(f'{created} créés')
        if updated:
            parts.append(f'{updated} mis à jour')
        if skipped:
            parts.append(f'{skipped} inchangés')
        if no_email:
            parts.append(f'{no_email} sans email ignorés')
        flash('Import vCard : ' + (', '.join(parts) or 'aucun contact'), 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erreur import vCard : {e}', 'error')
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass
    return redirect(url_for('contacts.index'))


@bp.route('/import', methods=['GET', 'POST'])
@admin_required
def index():
    listes = _active_listes()
    if request.method == 'POST':
        file = request.files.get('file')
        if not file or not file.filename:
            flash('Aucun fichier sélectionné', 'error')
            return redirect(url_for('imports.index'))
        fn = file.filename.lower()
        if fn.endswith('.vcf') or fn.endswith('.vcard'):
            return _import_vcard_direct(file)

        # Tabulaire (.xlsx / .csv / .tsv) → sauver en temp + écran de mapping
        os.makedirs(_IMPORT_DIR, exist_ok=True)
        _sweep_stale_imports()
        ext = '.xlsx' if fn.endswith(('.xlsx', '.xlsm')) else ('.tsv' if fn.endswith('.tsv') else '.csv')
        token = uuid.uuid4().hex
        path = os.path.join(_IMPORT_DIR, token + ext)
        file.save(path)
        try:
            headers, rows = _read_tabular(path, file.filename)
        except Exception as e:
            try:
                os.remove(path)
            except OSError:
                pass
            flash(f'Lecture du fichier impossible : {e}', 'error')
            return redirect(url_for('imports.index'))
        if not headers:
            try:
                os.remove(path)
            except OSError:
                pass
            flash("Fichier vide ou sans ligne d'en-tête.", 'error')
            return redirect(url_for('imports.index'))
        suggested = _suggest_mapping(headers, _import_targets())
        # Auto-application du meilleur mapping enregistré (≥2 colonnes reconnues),
        # par-dessus l'heuristique — modifiable ensuite.
        saved = _saved_mappings()
        best, best_n = _best_saved(headers, saved)
        applied = None
        mapping = suggested
        if best and best_n >= 2:
            mapping = dict(suggested)
            mapping.update(_apply_saved(headers, best))
            applied = best.name
        return _render_mapping(token, ext, file.filename, headers, rows, mapping,
                               request.form.get('list_id', ''),
                               request.form.get('new_list_name', ''),
                               request.form.get('update_existing') == 'on',
                               saved, applied_mapping=applied)

    return render_template('import.html', listes=listes)


@bp.route('/import/mapping', methods=['POST'])
@admin_required
def import_mapping():
    token = request.form.get('token', '')
    ext = request.form.get('ext', '')
    path = os.path.join(_IMPORT_DIR, token + ext)
    if not token or not os.path.exists(path):
        flash("Session d'import expirée — recommencez.", 'error')
        return redirect(url_for('imports.index'))

    filename = request.form.get('filename', 'x' + ext)
    mapping = _mapping_from_form(request.form)
    list_id = request.form.get('list_id', '')
    new_list_name = request.form.get('new_list_name', '')
    update_existing = request.form.get('update_existing') == 'on'
    action = request.form.get('action', 'preview')

    try:
        headers, rows = _read_tabular(path, filename)
    except Exception as e:
        flash(f'Lecture impossible : {e}', 'error')
        return redirect(url_for('imports.index'))

    saved_mappings = _saved_mappings()

    # Réappliquer un mapping enregistré (il prime sur les colonnes qu'il couvre ; le
    # reste des associations en cours est conservé). AVANT toute création de champ.
    if action == 'apply_saved':
        sm = (ImportMapping.query.get(request.form.get('saved_id', type=int))
              if request.form.get('saved_id') else None)
        if sm:
            mapping.update(_apply_saved(headers, sm))
            flash(f'Mapping « {sm.name} » appliqué.', 'success')
        else:
            flash('Mapping enregistré introuvable.', 'error')
        return _render_mapping(token, ext, filename, headers, rows, mapping, list_id,
                               new_list_name, update_existing, saved_mappings,
                               applied_mapping=(sm.name if sm else None))

    # Création éventuelle de champs perso demandés à l'import (colonnes « __new__ »).
    created_fields = _create_import_fields(mapping, request.form)
    if created_fields:
        flash('Champ personnalisé créé : ' + ', '.join(lbl for lbl, _ in created_fields), 'success')

    # Enregistrer l'association pour réutilisation (les « __new__ » viennent d'être
    # résolus en clés réelles → on ne sauvegarde que des clés concrètes).
    if action == 'save':
        name = (request.form.get('save_as') or '').strip()
        if not name:
            flash('Donnez un nom au mapping à enregistrer.', 'error')
        else:
            to_save = {h: k for h, k in mapping.items() if k and k != '__new__'}
            sm = ImportMapping.query.filter_by(name=name).first()
            if sm:
                sm.mapping = to_save
            else:
                db.session.add(ImportMapping(name=name, mapping=to_save,
                                             created_by_id=current_user.id))
            db.session.commit()
            flash(f'Mapping « {name} » enregistré ({len(to_save)} colonne(s)).', 'success')
            saved_mappings = _saved_mappings()
        return _render_mapping(token, ext, filename, headers, rows, mapping, list_id,
                               new_list_name, update_existing, saved_mappings)

    col_keys, custom_keys = _key_sets()
    custom_types = _custom_types()
    extra, redirect_list_id = _resolve_target_list(list_id, new_list_name)

    # Garde-fou : mettre à jour sans clé d'identification créerait des doublons
    # quasi-vides au lieu de retrouver les contacts. On BLOQUE l'import réel et on
    # renvoie l'utilisateur à l'écran (avertissement + message) au lieu de subir.
    blocked = update_existing and not _dedup_ok(mapping)
    if action == 'run' and blocked:
        flash("Mise à jour impossible sans clé d'identification : associez « Nom » et "
              "« Prénom » (ou l'UID) pour retrouver les contacts existants — ou décochez "
              "« Mettre à jour les contacts existants » pour créer de nouveaux contacts.", 'error')

    if action == 'run' and not blocked:
        try:
            counts = _run_import(rows, mapping, col_keys, custom_keys, custom_types, update_existing, extra, current_user.id)
        except Exception as e:
            db.session.rollback()
            flash(f'Erreur import : {e}', 'error')
            return redirect(url_for('imports.index'))
        try:
            os.remove(path)
        except OSError:
            pass
        msg = (f"Import terminé : {counts['created']} créés, {counts['updated']} mis à jour, "
               f"{counts['skipped']} inchangés")
        if counts.get('no_email'):
            msg += f" — dont {counts['no_email']} sans email (à compléter)"
        flash(msg + '.', 'success')
        if counts.get('backup'):
            import os as _os
            flash(f"Sauvegarde de sécurité créée avant l'import : {_os.path.basename(counts['backup'])}", 'info')
        # Retour sur la liste que l'on vient de peupler (sinon liste complète)
        if redirect_list_id:
            return redirect(url_for('contacts.index', liste=redirect_list_id))
        return redirect(url_for('contacts.index'))

    # « Tester l'import » (dry-run, rien écrit) → compteurs uniquement
    counts, _sample = _dry_run(rows, mapping, col_keys, custom_keys, custom_types, update_existing, extra)
    return _render_mapping(token, ext, filename, headers, rows, mapping, list_id,
                           new_list_name, update_existing, saved_mappings,
                           previewed=True, counts=counts)


@bp.route('/import/template')
@admin_required
def import_template():
    """Modèle CSV généré depuis le registre : en-têtes attendus + une ligne d'exemple."""
    targets = _import_targets()
    example_vals = {'email': 'jean.dupont@exemple.fr', 'nom': 'Dupont', 'prenom': 'Jean',
                    'civilite': 'Monsieur', 'telephone': '06 12 34 56 78',
                    'adresse_ville': 'Lodève', 'adresse_cp': '34700',
                    'listes': 'Sénatoriales 2026'}
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow([t['label'] for t in targets])
    w.writerow([example_vals.get(t['key'], '') for t in targets])
    return Response(out.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': 'attachment; filename=modele_import_contacts.csv'})


@bp.route('/export/vcard')
@admin_required
def export_vcard():
    from vcard_converter import create_vcard
    liste_id = request.args.get('liste', type=int)
    ids = request.args.get('ids', '').strip()
    version = request.args.get('version', '3.0')
    if version not in ('3.0', '4.0'):
        version = '3.0'

    if ids:
        id_list = [int(x) for x in ids.split(',') if x.strip().isdigit()]
        contacts = (Contact.query.filter(Contact.id.in_(id_list), Contact.is_deleted == False)
                    .order_by(Contact.nom, Contact.prenom).all())
        filename = 'contacts_selection.vcf'
    elif liste_id:
        liste = Liste.query.get_or_404(liste_id)
        contacts = liste.active_contacts
        filename = f'contacts_{liste.nom}.vcf'
    else:
        contacts = Contact.query.filter(Contact.is_deleted == False).order_by(Contact.nom, Contact.prenom).all()
        filename = 'contacts_all.vcf'

    lines = []
    for c in contacts:
        adr_parts = [c.adresse_rue, c.adresse_complement, c.adresse_ville,
                     c.adresse_cp, c.adresse_region, c.adresse_pays]
        adresse = ', '.join(p for p in adr_parts if p)
        row = {
            'UID': c.uid or '',
            'Nom, Prénom': f"{c.nom},{c.prenom}",
            'Nom Complet': f"{c.prenom} {c.nom}".strip(),
            'Email_Autre': c.email,
            'Tel_Cell': c.telephone or '',
            'Organisation': c.organisation or '',
            'Note': c.notes or '',
            'Adresse': adresse,
            'Catégories': ' | '.join(l.nom for l in c.listes),
        }
        vcard = create_vcard(row, version)
        lines.append(vcard.serialize())

    return Response(
        '\n'.join(lines),
        mimetype='text/vcard',
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )


# Anti-injection de formule (CSV/Excel) : une cellule commençant par un de ces
# caractères peut être interprétée comme une formule à l'ouverture dans Excel/LibreOffice.
# On préfixe d'un ESPACE → la cellule devient du texte inerte. L'espace est retiré au
# ré-import (les valeurs sont `.strip()`) → round-trip préservé. Corrige aussi
# l'affichage des tél. « +33… » (sinon vus comme une formule dans un CSV).
_CSV_FORMULA_CHARS = ('=', '+', '-', '@', '\t', '\r')
_XLSX_FORMULA_CHARS = ('=',)   # xlsx : seul « = » devient une formule (openpyxl) ; +/-/@ = strings sûrs


def _formula_guard(v, chars):
    s = '' if v is None else str(v)
    return ' ' + s if s[:1] in chars else s


# Colonnes d'export (partagées TSV / Excel) : (en-tête, extracteur).
_EXPORT_COLUMNS = [
    ('UID', lambda c: c.uid or ''),
    ('Nom', lambda c: c.nom or ''),
    ('Prenom', lambda c: c.prenom or ''),
    ('Genre', lambda c: c.genre or ''),
    ('Titre', lambda c: c.titre or ''),
    ('Email', lambda c: c.email or ''),
    ('Telephone', lambda c: c.telephone or ''),
    ('Organisation', lambda c: c.organisation or ''),
    ('Rue', lambda c: c.adresse_rue or ''),
    ('Complement', lambda c: c.adresse_complement or ''),
    ('Ville', lambda c: c.adresse_ville or ''),
    ('CP', lambda c: c.adresse_cp or ''),
    ('Region', lambda c: c.adresse_region or ''),
    ('Pays', lambda c: c.adresse_pays or ''),
    ('Source', lambda c: c.source or ''),
    ('Notes', lambda c: c.notes or ''),
    ('Listes', lambda c: ','.join(l.nom for l in c.listes)),
]


def _select_export_contacts(ids, liste_id):
    """(contacts, nom_de_fichier_sans_extension) selon sélection / liste / tous."""
    if ids:
        id_list = [int(x) for x in ids.split(',') if x.strip().isdigit()]
        contacts = (Contact.query.filter(Contact.id.in_(id_list), Contact.is_deleted == False)
                    .order_by(Contact.nom, Contact.prenom).all())
        return contacts, 'contacts_selection'
    if liste_id:
        liste = Liste.query.get_or_404(liste_id)
        return liste.active_contacts, f'contacts_{liste.nom}'
    contacts = Contact.query.filter(Contact.is_deleted == False).order_by(Contact.nom, Contact.prenom).all()
    return contacts, 'contacts_all'


@bp.route('/export')
@admin_required
def export_contacts():
    contacts, base = _select_export_contacts(request.args.get('ids', '').strip(),
                                             request.args.get('liste', type=int))
    fmt = request.args.get('format', 'tsv')
    headers = [h for h, _ in _EXPORT_COLUMNS]

    if fmt == 'xlsx':
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Contacts'
        ws.append(headers)
        for c in contacts:
            ws.append([_formula_guard(fn(c), _XLSX_FORMULA_CHARS) for _, fn in _EXPORT_COLUMNS])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Response(
            buf.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename={base}.xlsx'}
        )

    # TSV (par défaut)
    output = io.StringIO()
    writer = csv.writer(output, delimiter='\t')
    writer.writerow(headers)
    for c in contacts:
        writer.writerow([_formula_guard(fn(c), _CSV_FORMULA_CHARS) for _, fn in _EXPORT_COLUMNS])

    return Response(
        output.getvalue(),
        mimetype='text/tab-separated-values',
        headers={'Content-Disposition': f'attachment; filename={base}.tsv'}
    )
